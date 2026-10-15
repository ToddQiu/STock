"""Dated holdings imports; source weights are never rescaled to 100%."""

import base64
import binascii
import csv
import hashlib
import io
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException

from .research_ledger import ResearchLedger


SYMBOLS = ("IGV", "XSW", "WCLD", "BUG")
IDENTITY = "IGV-XSW-WCLD-BUG"
MAX_FILE_BYTES = 1_000_000
HEADERS = {
    "name": {"name", "company", "companyname", "holdingname", "securityname", "公司", "公司名称", "公司名稱", "公司名", "证券名称", "證券名稱", "持仓名称", "持倉名稱"},
    "symbol": {"ticker", "symbol", "code", "代码", "代碼", "股票代码", "股票代碼", "证券代码", "證券代碼"},
    "weight": {"weight", "weight%", "%weight", "weighting", "%ofnetassets", "权重", "權重", "权重%", "權重%", "占比", "占比%", "持仓权重%", "持倉權重%"},
    "type": {"assetclass", "type", "类型", "類型", "资产类别", "資產類別"},
}


def text(value):
    return "" if value is None else str(value).strip()


def holding_type(name, symbol, explicit=""):
    if explicit.upper() in {"EQUITY", "STOCK", "股票", "股權", "股权"}:
        return "stock"
    value = f"{name} {symbol} {explicit}".upper()
    if re.search(r"E-MINI|\bINDEX\b|期货|期貨", value) or explicit.upper() in {"FUTURE", "FUTURES"}:
        return "future"
    if re.search(r"CASH|现金|現金|CURRENCY|US DOLLAR|货币|貨幣|MONEY MARKET|MONEY FUND|TREASURY FUND|BLK CSH", value) or symbol.upper() in {"USD", "EUR", "GBP", "JPY", "HKD", "KRW"}:
        return "cash"
    if re.search(r"PAYABLE|RECEIVABLE|COLLATERAL|应收|應收|应付|應付|保证金|保證金|\bOTHER\b|其他", value):
        return "other"
    if explicit and explicit.upper() not in {"EQUITY", "STOCK", "股票", "股權", "股权"}:
        return "other"
    return "stock" if symbol else "other"


def holding_key(symbol, kind):
    if kind != "stock" or not symbol:
        return ""
    symbol = symbol.upper().strip()
    parts = symbol.rsplit(" ", 1)
    market = "US"
    if len(parts) == 2:
        market = {"US": "US", "JP": "JP", "KS": "KR", "LN": "GB", "HK": "HK", "CN": "CN", "CA": "CA"}.get(parts[1], "")
        symbol = parts[0]
        if not market:
            return ""
    else:
        for suffix, country in ((".KS", "KR"), (".KQ", "KR"), (".T", "JP"), (".L", "GB"), (".HK", "HK"), (".TO", "CA")):
            if symbol.endswith(suffix):
                market, symbol = country, symbol[:-len(suffix)]
                break
    if market == "US" and not re.fullmatch(r"[A-Z][A-Z0-9.\-]{0,14}", symbol):
        return ""
    if symbol in {"BRK.A", "BRK.B"}:
        symbol = symbol.replace(".", "-")
    return f"{market}:{symbol}"


def parse_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    match = re.search(r"(20\d{2})[年/.-](\d{1,2})[月/.-](\d{1,2})", text(value))
    if match:
        try:
            return date(*map(int, match.groups())).isoformat()
        except ValueError:
            pass
    match = re.search(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w* (\d{1,2}),? (20\d{2})", text(value), re.I)
    if match:
        try:
            return date(int(match[3]), ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(match[1].lower())+1, int(match[2])).isoformat()
        except ValueError:
            pass
    match = re.search(r"\b(\d{1,2})-(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-(20\d{2})", text(value), re.I)
    if match:
        try:
            return date(int(match[3]), ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(match[2].lower())+1, int(match[1])).isoformat()
        except ValueError:
            pass
    return ""


def parse_weight(value, number_format=""):
    if isinstance(value, bool) or value is None:
        raise ValueError("权重缺失或不是有效数字，未导入任何持仓")
    raw = text(value)
    try:
        number = float(raw.replace(",", "").replace("%", ""))
    except ValueError:
        raise ValueError(f"无效权重：{raw[:60]}") from None
    if not math.isfinite(number):
        raise ValueError("权重必须是有限数字")
    # Excel stores a percentage-formatted 0.01 as 1%; plain numbers are percent points.
    if isinstance(value, (int, float)) and "%" in number_format:
        number *= 100
    return number


def parse_table(rows):
    columns, start = {}, None
    for index, row in enumerate(rows[:50]):
        found = {}
        for col, cell in enumerate(row):
            header = re.sub(r"[\s()（）]", "", text(cell[0])).lower()
            for field, aliases in HEADERS.items():
                if header in aliases:
                    found[field] = col
        if all(key in found for key in ("name", "symbol", "weight")):
            columns, start = found, index + 1
            break
    if start is None:
        raise ValueError("找不到公司、代码、权重三列，请保留原始表头")
    holdings = []
    blank_separator = False
    for row_number, row in enumerate(rows[start:], start + 1):
        def cell(field):
            index = columns.get(field, len(row))
            return row[index] if index < len(row) else (None, "")
        name, symbol = text(cell("name")[0]), text(cell("symbol")[0])
        raw_weight, fmt = cell("weight")
        if not name and not symbol and raw_weight is None:
            blank_separator = True
            continue
        if blank_separator and name and not symbol and raw_weight is None and all(cell[0] is None for index, cell in enumerate(row) if index != columns["name"]):
            remaining = rows[row_number:]
            if all(not text(r[columns["symbol"]][0]) and r[columns["weight"]][0] is None for r in remaining if len(r) > max(columns["symbol"], columns["weight"])):
                break
        blank_separator = False
        if name in {"合计", "合計", "总计", "總計", "TOTAL", "Total"}:
            continue
        # Issuer CSV legal footers have one text cell, not a holding row.
        if not name and not symbol and raw_weight is not None and len(row) == 1:
            continue
        if not name:
            raise ValueError(f"第 {row_number} 行缺少持仓名称")
        kind = holding_type(name, symbol, text(cell("type")[0]))
        holdings.append({"name": name, "symbol": symbol, "weight": parse_weight(raw_weight, fmt),
                         "raw_weight": text(raw_weight), "number_format": fmt, "type": kind,
                         "key": holding_key(symbol, kind), "source_row": row_number})
    if not holdings or len(holdings) > 2000:
        raise ValueError("持仓行数必须为 1 至 2000 条")
    return holdings, start


def parse_import(data, filename, *, symbol="", as_of=""):
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("请选择不超过 1 MB 的 Excel 或 CSV")
    suffix = Path(filename).suffix.lower()
    tables, dates, sources = {}, {}, {}
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 500 or sum(info.file_size for info in archive.infolist()) > 20_000_000:
                    raise ValueError("Excel 解压内容过大")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                if len(workbook.worksheets) > 16:
                    raise ValueError("Excel 工作表过多")
                for sheet in workbook.worksheets:
                    rows = []
                    for index, row in enumerate(sheet.iter_rows()):
                        if index >= 2500 or len(row) > 40:
                            raise ValueError("Excel 的行数或列数过多")
                        rows.append([(cell.value, cell.number_format) for cell in row])
                    matched = [key for key in SYMBOLS if re.search(rf"\b{key}\b", sheet.title.upper())]
                    if not matched:
                        for row in rows[:20]:
                            if row and text(row[0][0]).strip(": ").lower() in {"ticker symbol", "fund ticker", "etf代码", "etf代碼"}:
                                matched = [key for key in SYMBOLS if key in [text(cell[0]).upper() for cell in row[1:]]]
                                break
                    if len(matched) == 1:
                        if matched[0] in tables:
                            raise ValueError(f"{matched[0]} 有多个持仓工作表，无法确定版本")
                        tables[matched[0]] = rows
                    else:
                        for row in rows:
                            values = [text(cell[0]).upper() for cell in row]
                            keys = [key for key in SYMBOLS if key in values]
                            row_dates = {parsed for cell in row if (parsed := parse_date(cell[0]))}
                            if len(keys) == 1 and len(row_dates) == 1:
                                dates[keys[0]] = row_dates.pop()
                            if len(keys) == 1:
                                sources[keys[0]] = [text(cell[0]) for cell in row if text(cell[0]).startswith(("https://", "http://"))]
            finally:
                workbook.close()
        except (BadZipFile, KeyError, OSError, ParseError, DefusedXmlException) as error:
            raise ValueError("无法读取 Excel 文件") from error
    elif suffix == ".csv":
        if symbol not in SYMBOLS:
            raise ValueError("请选择 CSV 对应的 ETF")
        try:
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        except UnicodeError:
            raise ValueError("CSV 请使用 UTF-8 编码") from None
        tables[symbol] = [[(cell, "") for cell in row] for row in rows]
    else:
        raise ValueError("仅支持 .xlsx 与 .csv")
    if not tables:
        raise ValueError("工作表名须包含 IGV、XSW、WCLD 或 BUG")
    if as_of and not parse_date(as_of):
        raise ValueError("持仓日期无效")
    funds = {}
    for key, rows in tables.items():
        holdings, start = parse_table(rows)
        metadata = rows[:start-1]
        dated_rows = [row for row in metadata if re.search(r"as of|as at|持仓|持倉|截至|^date$", " ".join(text(cell[0]) for cell in row), re.I)]
        candidates = {parsed for row in dated_rows or metadata for cell in row if (parsed := parse_date(cell[0]))}
        # Global X publishes its CSV date in MM/DD/YYYY format.
        for row in rows[:start-1]:
            for cell in row:
                match = re.search(r"as of (\d{2})/(\d{2})/(20\d{2})", text(cell[0]), re.I)
                if match:
                    month, day, year = map(int, match.groups())
                    candidates.add(date(year, month, day).isoformat())
        if len(candidates) > 1:
            raise ValueError(f"{key} 表内有多个日期，请保留明确的持仓日期")
        holding_date = next(iter(candidates), dates.get(key, ""))
        if holding_date and key in dates and dates[key] != holding_date:
            raise ValueError(f"{key} 表内日期与汇总日期不一致")
        if as_of and holding_date and parse_date(as_of) != holding_date:
            raise ValueError(f"{key} 输入日期与文件日期不一致")
        holding_date = holding_date or parse_date(as_of)
        if not holding_date:
            raise ValueError(f"{key} 缺少持仓日期，请补充日期后重新导入")
        funds[key] = {"symbol": key, "as_of": holding_date, "filename": Path(filename).name,
                      "source_sha256": hashlib.sha256(data).hexdigest(), "holdings": holdings,
                      "source_urls": sources.get(key, []),
                      "provenance": "用户导入；权重沿用文件，未重新归一化"}
    return funds


def save_holdings(ledger, funds):
    with ledger.session() as db:
        latest = db.execute("SELECT * FROM input_versions WHERE kind='etf_holdings' AND identity=? ORDER BY id DESC LIMIT 1", (IDENTITY,)).fetchone()
    previous = json.loads(latest["payload"]) if latest else {"funds": {}}
    merged = {**previous["funds"], **funds}
    saved = ledger.observe("etf_holdings", IDENTITY, {"funds": merged}, expected_version=latest["version"] if latest else 0)
    return {"imported": list(funds), "version": saved["version"], "input_id": saved["id"]}


def import_payload(ledger, payload):
    if not isinstance(payload.get("filename"), str) or not isinstance(payload.get("data"), str):
        raise ValueError("请选择持仓文件")
    try:
        data = base64.b64decode(payload["data"], validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("文件编码无效") from None
    return save_holdings(ledger, parse_import(data, payload["filename"], symbol=payload.get("symbol", ""), as_of=payload.get("as_of", "")))


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--database", type=Path, default=Path("data/research.db"))
    parser.add_argument("--symbol", default="")
    parser.add_argument("--as-of", default="")
    parser.add_argument("--source-url", default="")
    args = parser.parse_args()
    funds = parse_import(args.file.read_bytes(), args.file.name, symbol=args.symbol, as_of=args.as_of)
    if args.source_url:
        for fund in funds.values():
            fund.update(source_url=args.source_url, provenance="发行商官方下载；原始权重，未重新归一化")
    print(json.dumps(save_holdings(ResearchLedger(args.database), funds), ensure_ascii=False))
